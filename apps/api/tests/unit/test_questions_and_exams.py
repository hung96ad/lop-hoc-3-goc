from decimal import Decimal
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.db.base_class import Base
from app.db.session import engine


@pytest_asyncio.fixture(autouse=True)
async def init_db():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
        await conn.run_sync(Base.metadata.create_all)
    yield
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


@pytest.mark.asyncio
async def test_question_bank_and_exam_freeze_snapshot():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Register Teacher
        reg_res = await ac.post("/api/v1/auth/register", json={
            "email": "teacher_exam@example.com",
            "password": "Password123!",
            "full_name": "Thay Giao Exam",
            "role": "teacher"
        })
        token = reg_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 2. Create single_choice question
        sc_payload = {
            "type": "single_choice",
            "content": "Thủ đô của Việt Nam là gì?",
            "default_points": "2.00",
            "options": [
                {"content": "Hà Nội", "is_correct": True, "order_index": 0},
                {"content": "TP. Hồ Chí Minh", "is_correct": False, "order_index": 1},
                {"content": "Đà Nẵng", "is_correct": False, "order_index": 2}
            ]
        }
        sc_res = await ac.post("/api/v1/questions", json=sc_payload, headers=headers)
        assert sc_res.status_code == 200
        sc_q = sc_res.json()
        assert len(sc_q["options"]) == 3

        # 3. Create essay question with rubrics (3 Gốc - 3 Độc aligned)
        essay_payload = {
            "type": "essay",
            "content": "Trình bày hành động giúp đỡ bạn bè trong hoàn cảnh khó khăn.",
            "default_points": "8.00",
            "rubrics": [
                {
                    "criterion_name": "Biểu hiện Đạo đức (Vị tha, thấu cảm)",
                    "description": "Nhận thức về lòng trắc ẩn và tinh thần trách nhiệm",
                    "max_points": "4.00",
                    "level_descriptors": {"gioi": 4, "kha": 3, "trung_binh": 2}
                },
                {
                    "criterion_name": "Biểu hiện Trí tuệ & Nghị lực (Giải pháp thực tế)",
                    "description": "Biện pháp giải quyết vấn đề khách quan, kiên trì",
                    "max_points": "4.00",
                    "level_descriptors": {"gioi": 4, "kha": 3, "trung_binh": 2}
                }
            ]
        }
        es_res = await ac.post("/api/v1/questions", json=essay_payload, headers=headers)
        assert es_res.status_code == 200
        es_q = es_res.json()
        assert len(es_q["rubrics"]) == 2

        # 4. Duplicate question
        dup_res = await ac.post(f"/api/v1/questions/{sc_q['id']}/duplicate", headers=headers)
        assert dup_res.status_code == 200
        assert "[Bản sao]" in dup_res.json()["content"]

        # 5. Create Exam Draft with total 10 points
        exam_payload = {
            "title": "Đề thi Rèn luyện Đạo đức & Kiến thức",
            "description": "Đề kiểm tra định kỳ",
            "duration_minutes": 30,
            "total_score": "10.00",
            "questions": [
                {"question_id": sc_q["id"], "points": "2.00", "order_index": 0},
                {"question_id": es_q["id"], "points": "8.00", "order_index": 1}
            ]
        }
        exam_res = await ac.post("/api/v1/exams", json=exam_payload, headers=headers)
        assert exam_res.status_code == 200
        exam_data = exam_res.json()
        assert exam_data["status"] == "draft"
        assert len(exam_data["questions"]) == 2

        # Verify frozen snapshot
        frozen_q0 = exam_data["questions"][0]["frozen_question_data"]
        assert frozen_q0["content"] == "Thủ đô của Việt Nam là gì?"
        assert len(frozen_q0["options"]) == 3

        # 6. Publish Exam
        pub_res = await ac.post(f"/api/v1/exams/{exam_data['id']}/publish", headers=headers)
        assert pub_res.status_code == 200
        assert pub_res.json()["status"] == "published"


@pytest.mark.asyncio
async def test_generate_ai_rubric_three_roots():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Register Teacher
        reg_res = await ac.post("/api/v1/auth/register", json={
            "email": "teacher_rubric@example.com",
            "password": "Password123!",
            "full_name": "Thay Giao 3 Goc",
            "role": "teacher"
        })
        token = reg_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 2. Test generate rubric with empty content (should fail 400)
        fail_res = await ac.post("/api/v1/questions/generate-rubric", json={
            "content": "   ",
            "total_points": "10.00"
        }, headers=headers)
        assert fail_res.status_code == 400

        # 3. Test generate rubric with valid essay question
        gen_res = await ac.post("/api/v1/questions/generate-rubric", json={
            "content": "Khi đối mặt với thất bại trong kỳ thi, em sẽ vận dụng tinh thần Nghị lực và Trí tuệ như thế nào để vượt qua?",
            "total_points": "10.00"
        }, headers=headers)
        assert gen_res.status_code == 200
        rubrics = gen_res.json()["rubrics"]
        assert len(rubrics) == 3

        # Check 3 Gốc representations
        criterion_names = [r["criterion_name"] for r in rubrics]
        assert any("Đạo đức" in name for name in criterion_names)
        assert any("Trí tuệ" in name for name in criterion_names)
        assert any("Nghị lực" in name for name in criterion_names)

        # Check points sum matches requested total_points exactly
        total_pts = sum(Decimal(str(r["max_points"])) for r in rubrics)
        assert total_pts == Decimal("10.00")

        # 4. Create question directly using the AI-generated rubrics
        create_res = await ac.post("/api/v1/questions", json={
            "type": "essay",
            "content": "Khi đối mặt với thất bại trong kỳ thi, em sẽ vận dụng tinh thần Nghị lực và Trí tuệ như thế nào để vượt qua?",
            "default_points": "10.00",
            "rubrics": rubrics
        }, headers=headers)
        assert create_res.status_code == 200
        created_q = create_res.json()
        assert len(created_q["rubrics"]) == 3


@pytest.mark.asyncio
async def test_question_bank_managed_by_class():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Register Teacher
        reg_res = await ac.post("/api/v1/auth/register", json={
            "email": "teacher_class_q@example.com",
            "password": "Password123!",
            "full_name": "Thay Giao Lop Hoc",
            "role": "teacher"
        })
        token = reg_res.json()["access_token"]
        headers = {"Authorization": "Bearer " + token}

        # 2. Create 2 classes
        c1_res = await ac.post("/api/v1/classes", json={"name": "Lớp 10A1", "code": "10A1"}, headers=headers)
        assert c1_res.status_code == 200
        c1 = c1_res.json()

        c2_res = await ac.post("/api/v1/classes", json={"name": "Lớp 12 Chuyên Tin", "code": "12TIN"}, headers=headers)
        assert c2_res.status_code == 200
        c2 = c2_res.json()

        # 3. Create question assigned to Class 1
        q1_res = await ac.post("/api/v1/questions", json={
            "type": "single_choice",
            "content": "Câu hỏi dành riêng cho lớp 10A1",
            "default_points": "2.00",
            "class_id": c1["id"],
            "options": [
                {"content": "Đáp án 1", "is_correct": True, "order_index": 0},
                {"content": "Đáp án 2", "is_correct": False, "order_index": 1}
            ]
        }, headers=headers)
        assert q1_res.status_code == 200
        q1_data = q1_res.json()
        assert q1_data["class_id"] == c1["id"]
        assert q1_data["class_name"] == "Lớp 10A1"

        # 4. Create question assigned to Class 2
        q2_res = await ac.post("/api/v1/questions", json={
            "type": "single_choice",
            "content": "Câu hỏi dành riêng cho lớp 12 Tin",
            "default_points": "2.00",
            "class_id": c2["id"],
            "options": [
                {"content": "Đáp án A", "is_correct": True, "order_index": 0},
                {"content": "Đáp án B", "is_correct": False, "order_index": 1}
            ]
        }, headers=headers)
        assert q2_res.status_code == 200
        q2_data = q2_res.json()
        assert q2_data["class_id"] == c2["id"]
        assert q2_data["class_name"] == "Lớp 12 Chuyên Tin"

        # 5. Create general unassigned question
        q3_res = await ac.post("/api/v1/questions", json={
            "type": "single_choice",
            "content": "Câu hỏi ngân hàng chung",
            "default_points": "1.00",
            "class_id": None,
            "options": [
                {"content": "Đúng", "is_correct": True, "order_index": 0},
                {"content": "Sai", "is_correct": False, "order_index": 1}
            ]
        }, headers=headers)
        assert q3_res.status_code == 200
        q3_data = q3_res.json()
        assert q3_data["class_id"] is None
        assert q3_data["class_name"] is None

        # 6. Filter questions by Class 1
        filter_c1 = await ac.get(f"/api/v1/questions?class_id={c1['id']}", headers=headers)
        assert filter_c1.status_code == 200
        c1_items = filter_c1.json()
        assert len(c1_items) == 1
        assert c1_items[0]["id"] == q1_data["id"]

        # 7. Filter questions by Class 2
        filter_c2 = await ac.get(f"/api/v1/questions?class_id={c2['id']}", headers=headers)
        assert filter_c2.status_code == 200
        c2_items = filter_c2.json()
        assert len(c2_items) == 1
        assert c2_items[0]["id"] == q2_data["id"]

        # 8. Filter unassigned questions
        filter_unassigned = await ac.get("/api/v1/questions?class_id=unassigned", headers=headers)
        assert filter_unassigned.status_code == 200
        unassigned_items = filter_unassigned.json()
        assert len(unassigned_items) == 1
        assert unassigned_items[0]["id"] == q3_data["id"]

        # 9. List all questions
        all_res = await ac.get("/api/v1/questions", headers=headers)
        assert all_res.status_code == 200
        assert len(all_res.json()) == 3


@pytest.mark.asyncio
async def test_excel_template_and_bulk_import():
    import io
    import openpyxl

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Register Teacher
        reg_res = await ac.post("/api/v1/auth/register", json={
            "email": "teacher_excel@example.com",
            "password": "Password123!",
            "full_name": "Thay Giao Excel",
            "role": "teacher"
        })
        token = reg_res.json()["access_token"]
        headers = {"Authorization": "Bearer " + token}

        # Create class
        c_res = await ac.post("/api/v1/classes", json={"name": "Lớp 11 Văn", "code": "11VAN"}, headers=headers)
        class_id = c_res.json()["id"]

        # 2. Download template
        tpl_res = await ac.get("/api/v1/questions/import-template", headers=headers)
        assert tpl_res.status_code == 200
        assert "spreadsheetml.sheet" in tpl_res.headers["content-type"]
        wb = openpyxl.load_workbook(io.BytesIO(tpl_res.content))
        assert "Ngân hàng câu hỏi" in wb.sheetnames
        assert "Hướng dẫn nhập" in wb.sheetnames

        # 3. Create Excel buffer for bulk upload
        upload_wb = openpyxl.Workbook()
        ws = upload_wb.active
        ws.title = "Ngân hàng câu hỏi"
        ws.append([
            "STT", "Loại câu hỏi", "Nội dung câu hỏi", "Điểm", 
            "Đáp án A", "Đáp án B", "Đáp án C", "Đáp án D", 
            "Đáp án đúng", "Giải thích", "Rubric tự luận", "Mã lớp"
        ])
        # Row 1: Single choice
        ws.append([
            1, "single_choice", "Sông nào dài nhất Việt Nam?", 1.0,
            "Sông Đồng Nai", "Sông Hồng", "Sông Mê Kông", "Sông Cửu Long",
            "A", "Sông Đồng Nai là sông nội địa dài nhất VN", "", ""
        ])
        # Row 2: Multiple select
        ws.append([
            2, "multiple_select", "Chọn các số nguyên tố nhỏ hơn 10", 2.0,
            "2", "3", "4", "5",
            "A, B, D", "2, 3, 5 là các số nguyên tố", "", ""
        ])
        # Row 3: Essay without rubric (should auto-generate 3 roots)
        ws.append([
            3, "essay", "Hãy viết cảm nghĩ của em về lòng biết ơn đối với cha mẹ.", 10.0,
            "", "", "", "",
            "", "", "", ""
        ])
        # Row 4: Essay with custom rubric
        ws.append([
            4, "essay", "Phân tích tác phẩm Tắt đèn của Ngô Tất Tố.", 8.0,
            "", "", "", "",
            "", "", "Giá trị hiện thực: 4.0: Phản ánh chân thực | Nghệ thuật xây dựng nhân vật: 4.0: Khắc họa sinh động", ""
        ])

        excel_buf = io.BytesIO()
        upload_wb.save(excel_buf)
        excel_buf.seek(0)

        # 4. Upload file to /questions/import-excel
        files = {
            "file": ("test_questions.xlsx", excel_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        data = {
            "class_id": class_id
        }

        import_res = await ac.post("/api/v1/questions/import-excel", files=files, data=data, headers=headers)
        assert import_res.status_code == 200
        res_json = import_res.json()
        assert res_json["total"] == 4
        assert res_json["success_count"] == 4
        assert res_json["failed_count"] == 0
        assert len(res_json["imported_questions"]) == 4

        # 5. Verify all questions were saved and assigned to Lớp 11 Văn
        q_list_res = await ac.get(f"/api/v1/questions?class_id={class_id}", headers=headers)
        assert q_list_res.status_code == 200
        q_list = q_list_res.json()
        assert len(q_list) == 4

        # Verify essay question has 3 Gốc rubrics auto-generated
        essay_auto = next(q for q in q_list if q["content"] == "Hãy viết cảm nghĩ của em về lòng biết ơn đối với cha mẹ.")
        assert len(essay_auto["rubrics"]) == 3
        rubric_sum = sum(Decimal(str(r["max_points"])) for r in essay_auto["rubrics"])
        assert rubric_sum == Decimal("10.00")

        # Verify essay question with custom rubric
        essay_custom = next(q for q in q_list if "Tắt đèn" in q["content"])
        assert len(essay_custom["rubrics"]) == 2
        rubric_custom_sum = sum(Decimal(str(r["max_points"])) for r in essay_custom["rubrics"])
        assert rubric_custom_sum == Decimal("8.00")


@pytest.mark.asyncio
async def test_question_bank_managed_by_subject():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Register Teacher
        reg_res = await ac.post("/api/v1/auth/register", json={
            "email": "teacher_subject@example.com",
            "password": "Password123!",
            "full_name": "Thay Giao Mon Hoc",
            "role": "teacher"
        })
        token = reg_res.json()["access_token"]
        headers = {"Authorization": "Bearer " + token}

        # 2. Create Subjects (Chánh Kiến 1 & Nhân Tướng Học)
        s1_res = await ac.post("/api/v1/subjects", json={
            "name": "Chánh Kiến 1 - Đánh thức ý nghĩa cuộc đời",
            "code": "CK1",
            "description": "Khoá học nền tảng 3 Gốc Rễ của Thầy Trần Việt Quân"
        }, headers=headers)
        assert s1_res.status_code == 200
        s1 = s1_res.json()
        assert s1["code"] == "CK1"

        s2_res = await ac.post("/api/v1/subjects", json={
            "name": "Nhân Tướng Học ứng dụng",
            "code": "NTH",
            "description": "Nghệ thuật thấu hiểu con người"
        }, headers=headers)
        assert s2_res.status_code == 200
        s2 = s2_res.json()

        # 3. Create question assigned to Chánh Kiến 1
        q1_res = await ac.post("/api/v1/questions", json={
            "type": "single_choice",
            "content": "3 Gốc Rễ hình thành nên nhân cách cốt lõi là gì?",
            "default_points": "2.00",
            "subject_id": s1["id"],
            "options": [
                {"content": "Đạo đức - Trí tuệ - Nghị lực", "is_correct": True, "order_index": 0},
                {"content": "Tiền tài - Danh vọng - Quyền lực", "is_correct": False, "order_index": 1}
            ]
        }, headers=headers)
        assert q1_res.status_code == 200
        q1 = q1_res.json()
        assert q1["subject_id"] == s1["id"]
        assert q1["subject_name"] == "Chánh Kiến 1 - Đánh thức ý nghĩa cuộc đời"

        # 4. Create question assigned to Nhân Tướng Học
        q2_res = await ac.post("/api/v1/questions", json={
            "type": "single_choice",
            "content": "Tướng tự tâm sinh, câu nói này nhấn mạnh điều gì?",
            "default_points": "2.00",
            "subject_id": s2["id"],
            "options": [
                {"content": "Tâm thế nào thì lộ ra tướng mạo thế ấy", "is_correct": True, "order_index": 0},
                {"content": "Tướng mạo cố định không đổi", "is_correct": False, "order_index": 1}
            ]
        }, headers=headers)
        assert q2_res.status_code == 200

        # 5. Filter questions by subject CK1
        filter_s1 = await ac.get(f"/api/v1/questions?subject_id={s1['id']}", headers=headers)
        assert filter_s1.status_code == 200
        s1_items = filter_s1.json()
        assert len(s1_items) == 1
        assert s1_items[0]["content"] == "3 Gốc Rễ hình thành nên nhân cách cốt lõi là gì?"

        # 6. Verify subject question count in subjects list
        subjs_res = await ac.get("/api/v1/subjects", headers=headers)
        assert subjs_res.status_code == 200
        subjs = subjs_res.json()
        ck1_item = next(s for s in subjs if s["code"] == "CK1")
        assert ck1_item["question_count"] == 1


@pytest.mark.asyncio
async def test_import_chanh_kien_1_full_file():
    import os
    file_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "..", "40_cau_hoi_chanh_kien_1.xlsx"))
    assert os.path.exists(file_path), f"File {file_path} not found"

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Register Teacher
        reg_res = await ac.post("/api/v1/auth/register", json={
            "email": "teacher_ck1@example.com",
            "password": "Password123!",
            "full_name": "Thay Giao Chanh Kien",
            "role": "teacher"
        })
        token = reg_res.json()["access_token"]
        headers = {"Authorization": "Bearer " + token}

        # 2. Upload 40 questions Excel file
        with open(file_path, "rb") as f:
            file_bytes = f.read()

        files = {
            "file": ("40_cau_hoi_chanh_kien_1.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        res = await ac.post("/api/v1/questions/import-excel", files=files, headers=headers)
        assert res.status_code == 200
        data = res.json()
        assert data["total"] == 40
        assert data["success_count"] == 40
        assert data["failed_count"] == 0
        assert len(data["errors"]) == 0

        # 3. Verify all 40 questions are in database
        all_q_res = await ac.get("/api/v1/questions", headers=headers)
        assert all_q_res.status_code == 200
        all_q = all_q_res.json()
        assert len(all_q) == 40

        # 4. Verify subject CK1 exists and has 40 questions
        subjs_res = await ac.get("/api/v1/subjects", headers=headers)
        assert subjs_res.status_code == 200
        subjs = subjs_res.json()
        ck1 = next(s for s in subjs if s["code"] == "CK1")
        assert ck1["question_count"] == 40

        # 5. Filter by CK1 subject_id
        filtered_res = await ac.get(f"/api/v1/questions?subject_id={ck1['id']}", headers=headers)
        assert filtered_res.status_code == 200
        assert len(filtered_res.json()) == 40
