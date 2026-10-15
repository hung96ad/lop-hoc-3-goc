import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_ck1_excel(file_path: str):
    wb = openpyxl.Workbook()

    # --- Sheet 1: Ngân hàng câu hỏi ---
    ws = wb.active
    ws.title = "Ngân hàng câu hỏi"

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    headers = [
        "STT",
        "Loại câu hỏi",
        "Nội dung câu hỏi",
        "Điểm",
        "Đáp án A",
        "Đáp án B",
        "Đáp án C",
        "Đáp án D",
        "Đáp án đúng",
        "Giải thích",
        "Rubric tự luận",
        "Mã môn học",
        "Mã lớp"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    questions = [
        # --- CÂU 1 - 25: Trắc nghiệm 1 đáp án (single_choice) ---
        (
            1, "single_choice",
            "Cốt lõi của cây nhân cách con người theo triết lý 3 Gốc Rễ của Thầy Trần Việt Quân gồm những yếu tố nào?",
            2.00,
            "Đạo đức - Trí tuệ - Nghị lực",
            "Tiền tài - Danh vọng - Quyền lực",
            "Quan hệ - Ngoại hình - May mắn",
            "Kiến thức - Bằng cấp - Kỹ năng mềm",
            "A",
            "3 Gốc rễ Đạo đức - Trí tuệ - Nghị lực là nền tảng cốt lõi định hình nhân cách và sự thành công bền vững của con người.",
            "", "CK1", ""
        ),
        (
            2, "single_choice",
            "Định nghĩa cốt lõi của phẩm chất 'Đạo Đức' trong 3 Gốc là gì?",
            2.00,
            "Tâm vị tha lớn hơn tâm vị kỷ, luôn biết nghĩ và hành động vì cộng đồng",
            "Luôn tuân thủ tuyệt đối các quy định pháp luật vì sợ bị phạt",
            "Làm việc thiện để được mọi người xung quanh ca ngợi và vinh danh",
            "Chỉ giúp đỡ những người thân thiết có khả năng đền ơn cho mình",
            "A",
            "Đạo đức đích thực xuất phát từ tâm vị tha, lòng trắc ẩn và tinh thần trách nhiệm phụng sự vượt lên trên lợi ích vị kỷ cá nhân.",
            "", "CK1", ""
        ),
        (
            3, "single_choice",
            "Theo Chánh Kiến, 'Trí Tuệ' thực sự không phải là nhồi nhét nhiều thông tin mà là:",
            2.00,
            "Khả năng nhìn rõ bản chất thực tại, thấu suốt quy luật Nhân - Duyên - Quả",
            "Khả năng ghi nhớ siêu phàm mọi dữ liệu và có nhiều bằng cấp",
            "Sự khôn ngoan mưu mô để thu lợi về cho bản thân nhanh nhất",
            "Khả năng hùng biện hoạt ngôn và tranh cãi luôn giành chiến thắng",
            "A",
            "Trí tuệ là cái thấy biết như thật (như thị), hiểu rõ bản chất vấn đề và quy luật nhân duyên quả vận hành trong đời sống.",
            "", "CK1", ""
        ),
        (
            4, "single_choice",
            "Cột 'Nghị Lực' trong 3 Gốc được cấu thành bởi ba phẩm chất cốt lõi nào?",
            2.00,
            "Dũng - Nhẫn - Tĩnh",
            "Nhanh - Mạnh - Chuẩn",
            "Gan dạ - Liều lĩnh - Bất chấp",
            "Thông minh - Năng động - Quyết đoán",
            "A",
            "Nghị lực gồm 3 trụ cột: Dũng (dũng cảm đối diện sự thật), Nhẫn (kiên trì bền bỉ vượt khó) và Tĩnh (bình tĩnh làm chủ cảm xúc).",
            "", "CK1", ""
        ),
        (
            5, "single_choice",
            "Biểu hiện cao đẹp nhất của chữ 'Dũng' trong cột Nghị Lực là:",
            2.00,
            "Dũng cảm nhận lỗi, đối diện sự thật và dám làm việc khó vì lợi ích chung",
            "Liều mạng xông vào chỗ nguy hiểm mà không cần suy tính hậu quả",
            "Dám lớn tiếng tranh cãi bảo vệ cái tôi khi bị người khác phê bình",
            "Dám bất chấp quy tắc đạo đức để đạt được mục tiêu cá nhân",
            "A",
            "Dũng cảm đích thực là dũng khí của bậc trượng phu: dám nhận sai lầm, dám vượt qua nỗi sợ và dấn thân vì chân lý.",
            "", "CK1", ""
        ),
        (
            6, "single_choice",
            "Ý nghĩa chuẩn xác của chữ 'Nhẫn' trong cột Nghị Lực là gì?",
            2.00,
            "Kiên trì, bền bỉ vượt qua khó khăn gian khổ để hoàn thành mục tiêu đến cùng",
            "Cam chịu nhục nhã và bất công trong đau khổ mà không tìm giải pháp",
            "Chần chừ, trì hoãn và không dám dứt khoát đưa ra hành động",
            "Thờ ơ bỏ mặc mọi thứ xung quanh và trốn chạy thực tại",
            "A",
            "Nhẫn là sức mạnh của sự kiên định, bền bỉ, nhẫn nại trước chông gai nghịch cảnh chứ không phải sự hèn nhát cam chịu.",
            "", "CK1", ""
        ),
        (
            7, "single_choice",
            "Trạng thái 'Tĩnh' trong cột Nghị Lực được thể hiện như thế nào?",
            2.00,
            "Giữ được tâm an định, bình tĩnh, sáng suốt trước biến cố và áp lực nghịch cảnh",
            "Vô cảm, lạnh lùng, không quan tâm đến cảm xúc của bất kỳ ai",
            "Ngồi yên thụ động không làm gì cả để tránh mắc sai lầm",
            "Cách ly hoàn toàn khỏi xã hội để không phải đối mặt với khó khăn",
            "A",
            "Tĩnh là năng lực làm chủ tâm trí, giữ được sự tĩnh lặng và sáng suốt giữa giông bão cuộc đời.",
            "", "CK1", ""
        ),
        (
            8, "single_choice",
            "Trong 3 Độc (Tham - Sân - Si), độc nào là gốc rễ nguy hiểm nhất sinh ra hai độc còn lại?",
            2.00,
            "Si (Sự u mê, vô minh, thiếu hiểu biết đúng đắn về nhân quả)",
            "Tham (Ham muốn chiếm đoạt, ích kỷ, dính mắc)",
            "Sân (Sự tức giận, thù hận, bất toại nguyện)",
            "Mạn (Kiêu ngạo, tự cao tự đại)",
            "A",
            "Si (thiếu trí tuệ, vô minh) là gốc rễ sâu xa nhất. Vì si mê không hiểu rõ nhân quả nên mới khởi sinh lòng tham và tâm sân hận.",
            "", "CK1", ""
        ),
        (
            9, "single_choice",
            "Khi gặp chuyện bất như ý liền nổi giận, to tiếng quát mắng và đổ lỗi cho người khác, tâm nào đang chi phối mạnh nhất?",
            2.00,
            "Tâm Sân",
            "Tâm Tham",
            "Tâm Dũng",
            "Tâm Trí tuệ",
            "A",
            "Tâm Sân là phản ứng khó chịu, tức giận, chối bỏ khi thực tế không diễn ra theo ý muốn vị kỷ của bản thân.",
            "", "CK1", ""
        ),
        (
            10, "single_choice",
            "Hiện tượng một người luôn khao khát tích lũy của cải vô tận, so bì hơn thua và sợ mất mát bắt nguồn từ:",
            2.00,
            "Tâm Tham",
            "Tâm Sân",
            "Tâm Tĩnh",
            "Tâm Nghị lực",
            "A",
            "Tâm Tham vận hành theo cơ chế gom về cho mình, bám víu dính mắc vào của cải danh vọng và luôn thấy thiếu thốn bên trong.",
            "", "CK1", ""
        ),
        (
            11, "single_choice",
            "Một người có trí thông minh xuất chúng và nghị lực phi thường nhưng thiếu vắng Đạo Đức thì sẽ trở thành:",
            2.00,
            "Kẻ phá hoại nguy hiểm cho xã hội với quy mô lớn",
            "Bậc hiền triết dẫn dắt cộng đồng thoát nghèo",
            "Người công dân mẫu mực cống hiến thầm lặng",
            "Người bình thường không có ảnh hưởng gì tới ai",
            "A",
            "Tài năng và nghị lực khi không có đạo đức dẫn đường sẽ biến thành công cụ phục vụ lòng tham và gây tác hại to lớn cho nhân loại.",
            "", "CK1", ""
        ),
        (
            12, "single_choice",
            "Một người rất giàu lòng trắc ẩn thương người (Đạo đức) nhưng lại thiếu Trí tuệ và Nghị lực thì dễ mắc sai lầm nào?",
            2.00,
            "Thương hại mù quáng, lòng tốt bị lợi dụng và dễ bỏ cuộc khi gặp nghịch cảnh",
            "Xây dựng được các tổ chức từ thiện quy mô toàn cầu",
            "Giải quyết triệt để tận gốc rễ nỗi khổ của tha nhân",
            "Luôn kiên định vượt qua mọi sự phản đối của dư luận",
            "A",
            "Nhiệt tình cộng dốt nát thành phá hoại; đạo đức thiếu trí tuệ sẽ dẫn đến hành động sai lầm, thiếu nghị lực sẽ không làm việc thiện bền lâu.",
            "", "CK1", ""
        ),
        (
            13, "single_choice",
            "Việc cha mẹ quá nuông chiều con, đáp ứng mọi đòi hỏi và làm thay mọi việc thực chất đang làm suy yếu cột nào của con?",
            2.00,
            "Cột Nghị lực (khiến con yếu đuối, ỷ lại, sợ khó)",
            "Cột Chiều cao và thể lực",
            "Khả năng học thuộc lòng bài vở",
            "Cảm xúc lãng mạn nghệ thuật",
            "A",
            "Nuông chiều thái quá tước đoạt cơ hội rèn luyện ý chí vượt khó (Dũng - Nhẫn - Tĩnh), khiến con trở thành 'gà công nghiệp'.",
            "", "CK1", ""
        ),
        (
            14, "single_choice",
            "Khi mắc sai lầm nhưng dám can đảm thừa nhận trách nhiệm và sửa đổi thay vì chối bỏ, đó là biểu hiện của:",
            2.00,
            "Chữ 'Dũng' của Nghị Lực kết hợp cùng Đạo Đức trung thực",
            "Sự yếu đuối dễ bị người khác bắt nạt",
            "Mưu mẹo xoa dịu dư luận",
            "Tâm Sân ẩn ngầm",
            "A",
            "Dám nhìn thẳng vào khuyết điểm và nhận lỗi là biểu hiện tối cao của Dũng khí và Đạo đức trung thực với chính mình.",
            "", "CK1", ""
        ),
        (
            15, "single_choice",
            "Phương thuốc rèn luyện hàng ngày hữu hiệu nhất để từng bước chuyển hóa tâm Tham là gì?",
            2.00,
            "Thực tập bố thí, sẻ chia và cho đi vô điều kiện (tập buông xả)",
            "Tìm cách kiếm thêm thật nhiều tiền để thỏa mãn mọi ham muốn",
            "Chỉ trích và oán ghét bản thân mỗi khi khởi lòng muốn",
            "Đi mua sắm và du lịch xa xỉ để giải tỏa căng thẳng",
            "A",
            "Cách tốt nhất để trị bệnh gom vào (Tham) là tập thói quen mở rộng bàn tay sẻ chia (Bố thí, xả buông, phụng sự tha nhân).",
            "", "CK1", ""
        ),
        (
            16, "single_choice",
            "Công thức chuẩn xác về sự vận hành của quy luật Nhân Quả trong cuộc sống là:",
            2.00,
            "Nhân + Duyên = Quả",
            "Nhân + Quả = Số phận",
            "Duyên + May rủi = Thành công",
            "Hành động + Định mệnh = Kết quả",
            "A",
            "Mọi sự việc hiện tượng đều hình thành bởi hạt giống bên trong (Nhân) kết hợp cùng điều kiện môi trường bên ngoài (Duyên) tạo ra Kết quả (Quả).",
            "", "CK1", ""
        ),
        (
            17, "single_choice",
            "Trong công thức 'Nhân + Duyên = Quả', yếu tố 'Duyên' đóng vai trò bản chất là gì?",
            2.00,
            "Môi trường, hoàn cảnh bên ngoài tạo điều kiện cho hạt giống (Nhân) trổ hoa kết trái",
            "Số mệnh định sẵn bất biến do bề trên an bài",
            "Sự may rủi tình cờ không có bất kỳ quy luật nào",
            "Bản chất cốt lõi bên trong của hạt giống",
            "A",
            "Duyên là các yếu tố xúc tác như đất, nước, ánh sáng, phân bón hay môi trường giáo dục, bạn bè giúp nhân phát triển.",
            "", "CK1", ""
        ),
        (
            18, "single_choice",
            "Ba báu vật cuộc đời (Tam Bảo) theo Chánh Kiến giúp con người bồi đắp 3 Gốc Rễ gồm:",
            2.00,
            "Thầy hiền trí - Nhóm bạn tốt - Tủ sách tinh hoa",
            "Gia tài lớn - Địa vị cao - Danh tiếng vang dội",
            "Trường chuyên - Lớp chọn - Bằng thủ khoa",
            "Nhà lầu - Xe hơi - Quan hệ ngoại giao rộng",
            "A",
            "Tam bảo cuộc đời là môi trường sống quyết định vận mệnh: có Thầy dẫn đường, có Bạn đồng hành và có Sách tinh hoa nuôi dưỡng tâm trí.",
            "", "CK1", ""
        ),
        (
            19, "single_choice",
            "Thế nào là một người 'Thầy Hiền Trí' đích thực theo tinh thần Chánh Kiến?",
            2.00,
            "Người vừa có Trí tuệ sâu sắc dẫn lối, vừa có Đạo đức nhân cách sống gương mẫu",
            "Người dạy các mánh khóe làm giàu nhanh chóng không cần lao động chân chính",
            "Người có học hàm học vị cao nhất nhưng không bao giờ quan tâm học trò",
            "Người chỉ nói những điều êm tai nịnh nọt và không bao giờ nghiêm khắc",
            "A",
            "Thầy hiền trí là bậc minh sư có cả Tâm và Tầm: tâm vị tha yêu thương và tầm nhìn trí tuệ soi sáng con đường đúng đắn.",
            "", "CK1", ""
        ),
        (
            20, "single_choice",
            "'Tủ sách tinh hoa' khác biệt với các loại sách thông thường ở đặc tính nào?",
            2.00,
            "Đúc kết đạo lý cuộc đời của vĩ nhân, anh hùng, bậc giác ngộ và chân lý muôn đời",
            "Có giá bìa đắt đỏ nhất và được trang trí mạ vàng bóng bẩy",
            "Chứa các nội dung giật gân, ngôn tình lãng mạn theo xu hướng mạng",
            "Sách dày hàng ngàn trang với thuật ngữ rối rắm không áp dụng được",
            "A",
            "Sách tinh hoa là thức ăn cho tâm: đúc kết tinh túy tư tưởng vĩ nhân, bài học đạo lý nhân quả giúp khai sáng tâm trí con người.",
            "", "CK1", ""
        ),
        (
            21, "single_choice",
            "Đặc điểm nhận diện một 'Nhóm bạn tốt' trong Tam Bảo là:",
            2.00,
            "Cùng nhau rèn luyện 3 Gốc, dám chân thành chỉ ra lỗi sai để cùng tiến bộ",
            "Thường xuyên tụ tập ăn nhậu xa xỉ và luôn tán tụng hùa theo nhau",
            "Chỉ gặp nhau khi có cơ hội làm ăn vụ lợi cá nhân",
            "Không bao giờ dám góp ý vì sợ mất lòng nhau",
            "A",
            "Bạn tốt (thiện tri thức) là người đồng hành nâng đỡ lúc khó khăn, dám nói lời thẳng thắn chân tình vì sự tiến bộ của bạn.",
            "", "CK1", ""
        ),
        (
            22, "single_choice",
            "Định luật 'Hạt nhỏ cây to' trong quy luật Nhân Quả muốn nhắc nhở chúng ta điều gì?",
            2.00,
            "Một hành vi dù nhỏ bé, nếu lặp đi lặp lại qua thời gian sẽ tạo thành quả cực kỳ to lớn",
            "Việc nhỏ không quan trọng, chỉ cần chú ý làm việc lớn",
            "Quả luôn luôn xuất hiện tức thời ngay sau khi hành động",
            "Chỉ những người sinh ra trong gia đình lớn mới thành công",
            "A",
            "Một hạt sồi nhỏ có thể mọc thành cây sồi cổ thụ; một thói quen xấu nhỏ tích lũy thành đại họa, một việc thiện nhỏ thành đại phúc.",
            "", "CK1", ""
        ),
        (
            23, "single_choice",
            "Khi một người gieo nhân thiện nhưng đang gặp 'Nghịch duyên' (hoàn cảnh bất lợi), kết quả sẽ ra sao?",
            2.00,
            "Quả ngọt có thể chậm trổ, nhưng hạt giống thiện vẫn được bảo tồn và sẽ đơm hoa khi hội đủ duyên",
            "Hạt giống thiện sẽ bị tiêu diệt hoàn toàn và biến thành quả xấu",
            "Quy luật nhân quả bị vô hiệu hóa bởi hoàn cảnh nghịch lý",
            "Người đó nên từ bỏ làm việc thiện vì không có ích gì",
            "A",
            "Nhân chưa đủ duyên thì chưa trổ. Kiên trì giữ nhân lành, chuyển hóa duyên thì quả ngọt chắc chắn sẽ đến đúng thời điểm.",
            "", "CK1", ""
        ),
        (
            24, "single_choice",
            "Muốn thay đổi vận mệnh cuộc đời một cách bền vững và tận gốc rễ, con người cần bắt đầu từ:",
            2.00,
            "Chuyển hóa Nhân bên trong (rèn 3 Gốc) và Kiến tạo Duyên bên ngoài (môi trường Tam Bảo)",
            "Đi xem bói, đổi phong thủy và cầu khấn thần linh phù hộ",
            "Thay đổi diện mạo bên ngoài và sắm sửa đồ hiệu sang trọng",
            "Liên tục nhảy việc mỗi khi gặp áp lực công việc",
            "A",
            "Muốn đổi quả phải đổi nhân và duyên: tu sửa tâm tính (3 gốc) và đắm mình trong môi trường Tam bảo.",
            "", "CK1", ""
        ),
        (
            25, "single_choice",
            "Câu tục ngữ 'Gần mực thì đen, gần đèn thì rạng' khẳng định vai trò quyết định của yếu tố nào?",
            2.00,
            "Môi trường sống (Duyên), đặc biệt là sự ảnh hưởng của Thầy, Bạn và Sách",
            "Số phận định sẵn của mỗi cá nhân",
            "Tài sản thừa kế từ cha mẹ",
            "Sự may mắn ngẫu nhiên ngoài ý muốn",
            "A",
            "Môi trường sống có sức mạnh uốn nắn tâm thức rất lớn. Ở gần người có 3 gốc ta được soi sáng, ở gần 3 độc ta dễ bị lây nhiễm.",
            "", "CK1", ""
        ),

        # --- CÂU 26 - 35: Trắc nghiệm nhiều đáp án (multiple_select) ---
        (
            26, "multiple_select",
            "Những tiêu chí nào sau đây thể hiện bước 'Quan sát đa chiều' đúng đắn trong tiến trình Văn - Tư - Tu?",
            2.00,
            "Thu thập thông tin khách quan từ nhiều góc độ và nhân vật khác nhau",
            "Quan sát sâu sắc diễn biến tâm lý và bối cảnh thực tế mà không phán xét vội vàng",
            "Đưa ra kết luận ngay lập tức dựa trên cảm xúc yêu ghét cá nhân",
            "Giữ tâm khách quan, tỉnh táo, không để thiên kiến che mờ sự thật",
            "A, B, D",
            "Quan sát đa chiều đòi hỏi nhìn nhận đa diện, khách quan, tỉnh thức và không vội vã phán xét quy chụp.",
            "", "CK1", ""
        ),
        (
            27, "multiple_select",
            "Trong bước 'Phân tích chiều sâu' của tiến trình tư duy Chánh Kiến, chúng ta cần tìm kiếm điều gì?",
            2.00,
            "Tách lớp vấn đề và truy tìm chuỗi Nhân - Duyên - Quả cốt lõi bên dưới",
            "Tìm kiếm đối tượng cụ thể để đổ lỗi và trút giận",
            "Nhận diện quy luật vận hành và nguyên lý ngầm chi phối hiện tượng",
            "Phân biệt rạch ròi đâu là cốt lõi bản chất, đâu là lớp vỏ bề ngoài",
            "A, C, D",
            "Phân tích chiều sâu là bóc tách từng lớp vấn đề để thấy rõ quy luật nhân duyên và nguyên nhân gốc rễ chứ không tìm người đổ lỗi.",
            "", "CK1", ""
        ),
        (
            28, "multiple_select",
            "Một bản 'Đúc kết cốt lõi' có giá trị thực tiễn cao thường mang những đặc điểm nào?",
            2.00,
            "Đơn giản, ngắn gọn, súc tích (có thể sơ đồ hóa trực quan)",
            "Chỉ ra được nguyên lý nhân quả hoặc chân lý vận hành cốt lõi",
            "Có tính ứng dụng cao, dễ dàng chuyển hóa thành hành động thực tế",
            "Càng dài dòng hàng chục trang và chứa nhiều từ ngữ trừu tượng càng tốt",
            "A, B, C",
            "Đúc kết cốt lõi phải đơn giản hóa điều phức tạp, chạm tới chân lý và soi đường cho hành động thực tế mỗi ngày.",
            "", "CK1", ""
        ),
        (
            29, "multiple_select",
            "Theo Thầy Trần Việt Quân, mô hình '4 Vòng tròn Đào tạo' gồm những tầng nào (tính từ ngoài vào trong)?",
            2.00,
            "Vòng 1 (ngoài cùng): Kiến thức (Thông tin lý thuyết, văn bằng chứng chỉ)",
            "Vòng 2: Kỹ năng (Kỹ năng giao tiếp, chuyên môn, quản lý công việc)",
            "Vòng 3: Xu hướng tính cách & Năng khiếu sở trường (DISC, MBTI,...)",
            "Vòng 4 (hạt nhân trung tâm): Nhân cách cốt lõi 3 Gốc (Đạo đức - Trí tuệ - Nghị lực)",
            "A, B, C, D",
            "4 Vòng tròn đào tạo định vị chuẩn mực phát triển con người toàn diện, lấy nhân cách 3 Gốc làm lõi hạt nhân vững chắc.",
            "", "CK1", ""
        ),
        (
            30, "multiple_select",
            "Những căn bệnh phổ biến của nền giáo dục chạy theo bệnh thành tích hiện nay là gì?",
            2.00,
            "Quá chú trọng nhồi nhét Kiến thức Vòng 1 mà bỏ quên Nhân cách cốt lõi Vòng 4",
            "Đánh giá giá trị một con người chỉ qua điểm số và văn bằng giấy tờ",
            "Bao bọc nuông chiều khiến học sinh thiếu Nghị lực vượt khó (thiếu Dũng - Nhẫn - Tĩnh)",
            "Tập trung rèn luyện học sinh lối sống tự lập, biết ơn và phụng sự xã hội",
            "A, B, C",
            "Giáo dục nhồi sọ chạy theo điểm số tạo ra những con người bằng cấp cao nhưng tâm lý mỏng manh, thiếu nghị lực và thiếu định hướng sống.",
            "", "CK1", ""
        ),
        (
            31, "multiple_select",
            "Những hành vi nào dưới đây thể hiện 'Tâm Sân' đang vận hành tinh vi và ngấm ngầm?",
            2.00,
            "Bằng mặt nhưng không bằng lòng, ghim giữ ấm ức và thù dai trong lòng",
            "Thở dài, tỏ thái độ khó chịu hoặc mỉa mai khi người khác làm trái ý mình",
            "Chân thành cảm ơn khi được bạn bè chỉ ra điểm yếu và khuyết điểm",
            "Nói xấu sau lưng, châm chọc đối phương khi tranh luận bất thành",
            "A, B, D",
            "Tâm Sân không chỉ là cơn thịnh nộ bùng phát mà còn ẩn nấp dưới dạng chê bai ngầm, mỉa mai, ấm ức, thở dài, bất toại nguyện.",
            "", "CK1", ""
        ),
        (
            32, "multiple_select",
            "Những thói quen nào sau đây giúp bồi dưỡng và phát triển cột 'Trí Tuệ' mỗi ngày?",
            2.00,
            "Đọc sách tinh hoa và đúc kết lại bài học ứng dụng thực tế",
            "Trước mỗi quyết định quan trọng luôn tự vấn: 'Nhân này gieo ra sẽ trổ quả gì?'",
            "Lắng nghe góc nhìn phản biện đa chiều và sẵn sàng học hỏi cái đúng",
            "Tiếp nhận mọi thông tin giật gân trên mạng xã hội mà không cần tư duy phân tích",
            "A, B, C",
            "Rèn trí tuệ qua việc học sách tinh hoa, tư duy nhân quả sâu sắc và luôn giữ tinh thần cầu thị phản biện khách quan.",
            "", "CK1", ""
        ),
        (
            33, "multiple_select",
            "Để xây dựng một gia đình hạnh phúc bền vững theo 3 Gốc, cha mẹ cần thực hành điều gì?",
            2.00,
            "Cha mẹ làm gương trước (Thân giáo) thay vì chỉ ra lệnh bằng lời nói (Khẩu giáo)",
            "Tạo cơ hội cho con rèn luyện tính tự lập và nghị lực vượt khó từ việc nhà",
            "Nuông chiều bảo bọc con hết mức, làm thay mọi việc để con chuyên tâm học sách vở",
            "Kiến tạo tủ sách tinh hoa và môi trường Tam Bảo ngay trong không gian gia đình",
            "A, B, D",
            "Thân giáo là bài học sâu sắc nhất cho con. Gia đình có Tam bảo và rèn nghị lực cho con từ bé sẽ tạo ra những thế hệ vững vàng.",
            "", "CK1", ""
        ),
        (
            34, "multiple_select",
            "Một người lãnh đạo dẫn dắt doanh nghiệp theo triết lý 3 Gốc sẽ có những biểu hiện nào?",
            2.00,
            "Dám nhận lỗi về mình trước tiên khi xảy ra sự cố (Đạo đức & Dũng)",
            "Chăm lo bồi dưỡng, nâng đỡ cấp dưới phát triển cả nhân cách lẫn chuyên môn (Tâm phụng sự)",
            "Bóc lột cạn kiệt sức khỏe nhân viên để tối đa hóa lợi nhuận cá nhân",
            "Giữ tâm bình tĩnh, kiên định với sứ mệnh mang lại giá trị bền vững cho xã hội",
            "A, B, D",
            "Lãnh đạo 3 Gốc lấy phụng sự và đạo đức làm tôn chỉ, dùng trí tuệ định hướng và nghị lực dẫn dắt tập thể vượt qua sóng gió.",
            "", "CK1", ""
        ),
        (
            35, "multiple_select",
            "Những dấu hiệu nào cho thấy một người đang từng bước chiến thắng tâm 'Si' (Vô minh)?",
            2.00,
            "Nhìn nhận sự thật như nó đang là, không tự lừa dối mình và không phóng đại cảm xúc",
            "Hiểu sâu sắc quy luật Vô thường và Nhân quả, bớt dính mắc mê mờ vào ảo tưởng",
            "Nhận biết rõ ràng từng cơn sóng cảm xúc (tham, giận, sợ hãi) khi chúng vừa khởi phát trong tâm",
            "Tìm đến bói toán, cúng sao giải hạn để cầu mong giàu sang bình an tức thì",
            "A, B, C",
            "Vượt qua tâm Si là bước vào vùng sáng của Chánh Kiến: thấy rõ nhân duyên quả, thấu tỏ vô thường và luôn có năng lực tự soi tâm mình.",
            "", "CK1", ""
        ),

        # --- CÂU 36 - 40: Câu hỏi Tự luận (essay) tích hợp Rubric 3 Gốc ---
        (
            36, "essay",
            "Phân tích tiến trình tư duy 'Quan sát - Phân tích - Đúc kết' (Văn - Tư - Tu). Hãy vận dụng tiến trình này để giải quyết một thói quen xấu hoặc một bế tắc cụ thể trong công việc/cuộc sống của em.",
            10.00,
            "", "", "", "", "",
            "Bài làm cần nêu rõ định nghĩa từng bước và bài học vận dụng thực tế sâu sắc.",
            "Đạo đức (Tâm vị tha & Trách nhiệm): 3.50: Nhận diện trung thực thói quen xấu của bản thân, không né tránh đổ lỗi; thể hiện trách nhiệm cải thiện để không gây ảnh hưởng tiêu cực đến người xung quanh. | Trí tuệ (Tư duy Nhân Quả & Giải pháp): 3.50: Phân tích nguồn gốc sâu xa của vấn đề theo luật Nhân - Duyên - Quả; đúc kết thành nguyên tắc hành động ngắn gọn, khả thi. | Nghị lực (Dũng - Nhẫn - Tĩnh): 3.00: Dũng cảm cam kết từ bỏ thói quen xấu, lên kế hoạch kỷ luật kiên trì thực hiện đều đặn mỗi ngày.",
            "CK1", ""
        ),
        (
            37, "essay",
            "Thế nào là 'Tam Bảo' cuộc đời (Thầy hiền trí - Tủ sách tinh hoa - Nhóm bạn tốt)? Hãy đánh giá môi trường sống hiện tại của em theo mô hình Tam Bảo và đề xuất kế hoạch cụ thể để kiến tạo môi trường Tam Bảo cho bản thân trong 1 năm tới.",
            10.00,
            "", "", "", "", "",
            "Bài làm cần phân tích vai trò của từng báu vật và đưa ra kế hoạch hành động chi tiết.",
            "Đạo đức (Tâm vị tha & Trách nhiệm): 3.50: Thấu hiểu giá trị của môi trường thiện lành, có ý thức trở thành bạn tốt để nâng đỡ người khác cùng tiến bộ. | Trí tuệ (Tư duy Nhân Quả & Giải pháp): 3.50: Đánh giá khách quan, trung thực hiện trạng môi trường của mình; đề xuất tiêu chí chọn Thầy, chọn Sách, chọn Bạn sâu sắc, đúng đắn. | Nghị lực (Dũng - Nhẫn - Tĩnh): 3.00: Quyết đoán dũng cảm cắt giảm các mối quan hệ độc hại, kiên trì duy trì kỷ luật đọc sách và sinh hoạt nhóm bạn tốt hàng tuần.",
            "CK1", ""
        ),
        (
            38, "essay",
            "Trình bày cách nhận diện 3 Độc (Tham - Sân - Si) trong các mối quan hệ gia đình hoặc công sở. Hãy phân tích một tình huống xung đột điển hình và đưa ra giải pháp chuyển hóa 3 Độc thành 3 Gốc (Đạo đức - Trí tuệ - Nghị lực).",
            10.00,
            "", "", "", "", "",
            "Bài làm cần nêu rõ cơ chế vận hành của 3 độc và phương pháp chuyển hóa nhân duyên quả thực tế.",
            "Đạo đức (Tâm vị tha & Trách nhiệm): 3.50: Lấy tâm từ bi, thấu cảm và lợi ích tập thể làm trọng; sẵn sàng buông bỏ cái tôi ích kỷ để bao dung tha thứ. | Trí tuệ (Tư duy Nhân Quả & Giải pháp): 3.50: Chỉ rõ tâm Tham, Sân hay Si đang giật dây xung đột; đưa ra giải pháp thấu tình đạt lý, giải quyết tận gốc rễ nhân duyên. | Nghị lực (Dũng - Nhẫn - Tĩnh): 3.00: Giữ tâm Tĩnh lặng không phản ứng giận dữ tức thì; Nhẫn nại lắng nghe đối phương và Dũng cảm nhận phần lỗi của mình.",
            "CK1", ""
        ),
        (
            39, "essay",
            "Phân tích mô hình '4 Vòng tròn Đào tạo' của Thầy Trần Việt Quân. Tại sao Nhân cách cốt lõi 3 Gốc (Vòng 4) lại là hạt nhân quyết định sự thành bại và hạnh phúc bền vững của một đời người hơn cả Kiến thức (Vòng 1) và Kỹ năng (Vòng 2)?",
            10.00,
            "", "", "", "", "",
            "Bài làm cần phân tích rõ thứ tự ưu tiên của 4 vòng tròn và minh họa bằng dẫn chứng thực tế thuyết phục.",
            "Đạo đức (Tâm vị tha & Trách nhiệm): 3.50: Hiểu rõ mục đích học tập và làm việc là để phụng sự cuộc đời và làm việc có ích, chứ không chỉ vị kỷ tích lũy danh lợi. | Trí tuệ (Tư duy Nhân Quả & Giải pháp): 3.50: Lập luận logic, dẫn chứng thực tế thuyết phục về người có tài nhưng thiếu đức hoặc người có 3 gốc vững vàng thành tựu rực rỡ. | Nghị lực (Dũng - Nhẫn - Tĩnh): 3.00: Kiên định với con đường tu dưỡng bên trong, không chạy theo áp lực hào nhoáng hình thức bên ngoài.",
            "CK1", ""
        ),
        (
            40, "essay",
            "Khi đứng trước một biến cố lớn hoặc một nghịch cảnh thất bại đau đớn trong cuộc sống, một người có Nghị lực (Dũng - Nhẫn - Tĩnh) và hiểu biết sâu sắc về Nhân Quả sẽ ứng xử như thế nào để vượt qua và chuyển hóa nghịch cảnh thành bài học vô giá?",
            10.00,
            "", "", "", "", "",
            "Bài làm thể hiện tư duy nhân quả sâu sắc và ý chí kiên định làm chủ số phận trước giông bão cuộc đời.",
            "Đạo đức (Tâm vị tha & Trách nhiệm): 3.50: Không oán trời trách người, nhận trọn vẹn trách nhiệm về mình; giữ được lòng vị tha và tình yêu thương trong nghịch cảnh. | Trí tuệ (Tư duy Nhân Quả & Giải pháp): 3.50: Nhìn thấy nghịch cảnh là tấm gương phản chiếu nhân quả và là đề thi cuộc đời để rèn luyện; tìm ra bài học sâu sắc nhất từ thất bại. | Nghị lực (Dũng - Nhẫn - Tĩnh): 3.00: Thể hiện trọn vẹn Dũng (dám đối diện), Nhẫn (kiên cường không bỏ cuộc), Tĩnh (an nhiên, không hoảng loạn tuyệt vọng).",
            "CK1", ""
        )
    ]

    data_font = Font(name="Arial", size=10)
    for row_idx, row_data in enumerate(questions, start=2):
        ws.append(list(row_data))
        ws.row_dimensions[row_idx].height = 28
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    col_widths = {
        1: 8,   # STT
        2: 18,  # Loại câu hỏi
        3: 50,  # Nội dung
        4: 10,  # Điểm
        5: 25,  # Đáp án A
        6: 25,  # Đáp án B
        7: 25,  # Đáp án C
        8: 25,  # Đáp án D
        9: 15,  # Đáp án đúng
        10: 30, # Giải thích
        11: 45, # Rubric
        12: 15, # Mã môn học
        13: 14  # Mã lớp
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # --- Sheet 2: Hướng dẫn nhập ---
    ws_guide = wb.create_sheet(title="Hướng dẫn nhập")
    guide_headers = ["Tên cột", "Bắt buộc", "Định dạng / Giá trị chấp nhận", "Ghi chú & Ví dụ"]
    ws_guide.append(guide_headers)
    for col_idx in range(1, len(guide_headers) + 1):
        cell = ws_guide.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_guide.row_dimensions[1].height = 26

    guides = [
        ("Loại câu hỏi", "Có", "single_choice / multiple_select / essay", "single_choice: 1 đáp án đúng\nmultiple_select: nhiều đáp án đúng\nessay: tự luận"),
        ("Nội dung câu hỏi", "Có", "Văn bản nội dung câu hỏi", "Nhập rõ ràng câu hỏi cần kiểm tra"),
        ("Điểm", "Có", "Số thực dương (VD: 2.0, 10.0)", "Mặc định 2.0 cho trắc nghiệm, 10.0 cho tự luận"),
        ("Đáp án A, B, C, D", "Tuỳ loại", "Văn bản các lựa chọn", "Bắt buộc tối thiểu 2 đáp án đối với câu hỏi trắc nghiệm"),
        ("Đáp án đúng", "Tuỳ loại", "A, B, C, D (hoặc A, B hoặc A, B, D)", "Đối với single_choice: chỉ 1 đáp án (VD: A)\nĐối với multiple_select: các đáp án cách nhau dấu phẩy (VD: A, B, D)"),
        ("Giải thích", "Không", "Văn bản giải thích đáp án", "Hiển thị cho học viên sau khi hoàn thành bài thi"),
        ("Rubric tự luận", "Không", "Tên tiêu chí: Điểm: Mô tả | Tên tiêu chí: Điểm: Mô tả", "Dành cho câu tự luận. Nếu để trống, hệ thống sẽ tự động tạo bộ Rubric 3 Gốc chuẩn (Đạo đức 35%, Trí tuệ 35%, Nghị lực 30%) khớp với điểm câu hỏi."),
        ("Mã môn học", "Không", "Mã code của môn học (VD: CK1)", "Môn học mà câu hỏi này trực thuộc. Nếu chưa có môn, hệ thống sẽ tự động tạo mới."),
        ("Mã lớp", "Không", "Mã code của lớp học (VD: LOP10A1)", "Tùy chọn nếu muốn phân định thêm theo lớp.")
    ]

    for r_idx, (cname, req, val, note) in enumerate(guides, start=2):
        ws_guide.append([cname, req, val, note])
        ws_guide.row_dimensions[r_idx].height = 36
        for c_idx in range(1, 5):
            cell = ws_guide.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws_guide.column_dimensions["A"].width = 20
    ws_guide.column_dimensions["B"].width = 12
    ws_guide.column_dimensions["C"].width = 40
    ws_guide.column_dimensions["D"].width = 50

    wb.save(file_path)
    print(f"File created successfully at {file_path}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "40_cau_hoi_chanh_kien_1.xlsx"
    build_ck1_excel(target)
